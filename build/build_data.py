#!/usr/bin/env python3
"""帰宅クイズ data.js 재생성.
입력:
  --jp   「EJU 일본 지리·연표 자료집.docx」 (수업/03_실전반/3_자료집/B_지도)
  --db   EJU_연표DB.xlsx (수업/03_실전반/3_자료집/A_연표)
  --wg   「EJU 세계지리 자료집.docx」 (같은 B_지도 폴더) — 자연지리·주요 통계
  build/concept/20xx-x.json  개념 문항 (SPEC.md 규격)
출력: ../data.js
필요 패키지: python-docx openpyxl japanmap numpy
"""
import argparse, json, re, glob, os, collections, unicodedata, sys
import docx, openpyxl, numpy as np, japanmap as jm
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
HERE = os.path.dirname(os.path.abspath(__file__))

def celltext(tc):
    """ruby(후리가나) 읽기(w:rt)는 빼고 본문만"""
    out = []
    for p in tc.iter(W+'p'):
        s = ''
        for el in p.iter():
            if el.tag == W+'t':
                anc, skip = el.getparent(), False
                while anc is not None and anc is not p:
                    if anc.tag == W+'rt': skip = True; break
                    anc = anc.getparent()
                if not skip: s += el.text or ''
        out.append(s.strip())
    return ' / '.join(x for x in out if x)

PREF_TXT = """1|北海道|홋카이도|札幌|삿포로|0
2|青森県|아오모리현|青森|아오모리|1
3|岩手県|이와테현|盛岡|모리오카|1
4|宮城県|미야기현|仙台|센다이|1
5|秋田県|아키타현|秋田|아키타|1
6|山形県|야마가타현|山形|야마가타|1
7|福島県|후쿠시마현|福島|후쿠시마|1
8|茨城県|이바라키현|水戸|미토|2
9|栃木県|도치기현|宇都宮|우쓰노미야|2
10|群馬県|군마현|前橋|마에바시|2
11|埼玉県|사이타마현|さいたま|사이타마|2
12|千葉県|지바현|千葉|지바|2
13|東京都|도쿄도|東京|도쿄|2
14|神奈川県|가나가와현|横浜|요코하마|2
15|新潟県|니가타현|新潟|니가타|3
16|富山県|도야마현|富山|도야마|3
17|石川県|이시카와현|金沢|가나자와|3
18|福井県|후쿠이현|福井|후쿠이|3
19|山梨県|야마나시현|甲府|고후|3
20|長野県|나가노현|長野|나가노|3
21|岐阜県|기후현|岐阜|기후|3
22|静岡県|시즈오카현|静岡|시즈오카|3
23|愛知県|아이치현|名古屋|나고야|3
24|三重県|미에현|津|쓰|4
25|滋賀県|시가현|大津|오쓰|4
26|京都府|교토부|京都|교토|4
27|大阪府|오사카부|大阪|오사카|4
28|兵庫県|효고현|神戸|고베|4
29|奈良県|나라현|奈良|나라|4
30|和歌山県|와카야마현|和歌山|와카야마|4
31|鳥取県|돗토리현|鳥取|돗토리|5
32|島根県|시마네현|松江|마쓰에|5
33|岡山県|오카야마현|岡山|오카야마|5
34|広島県|히로시마현|広島|히로시마|5
35|山口県|야마구치현|山口|야마구치|5
36|徳島県|도쿠시마현|徳島|도쿠시마|6
37|香川県|가가와현|高松|다카마쓰|6
38|愛媛県|에히메현|松山|마쓰야마|6
39|高知県|고치현|高知|고치|6
40|福岡県|후쿠오카현|福岡|후쿠오카|7
41|佐賀県|사가현|佐賀|사가|7
42|長崎県|나가사키현|長崎|나가사키|7
43|熊本県|구마모토현|熊本|구마모토|7
44|大分県|오이타현|大分|오이타|7
45|宮崎県|미야자키현|宮崎|미야자키|7
46|鹿児島県|가고시마현|鹿児島|가고시마|7
47|沖縄県|오키나와현|那覇|나하|7"""
REGIONS = [dict(ja='北海道地方',ko='홋카이도 지방'),dict(ja='東北地方',ko='도호쿠 지방'),dict(ja='関東地方',ko='간토 지방'),dict(ja='中部地方',ko='주부 지방'),dict(ja='近畿地方',ko='긴키 지방'),dict(ja='中国地方',ko='주고쿠 지방'),dict(ja='四国地方',ko='시코쿠 지방'),dict(ja='九州地方',ko='규슈 지방(오키나와 포함)')]

def build_prefs(d):
    kwtab = None
    for t in d.tables:
        try: h = [c.text.strip() for c in t.rows[0].cells]
        except Exception: continue
        if h[:2] == ['No.', '都道府県 · 県庁'] and len(t.rows) == 48: kwtab = t; break
    kw = {}
    if kwtab:
        for r in kwtab.rows[1:]:
            c = [x.text.strip() for x in r.cells]
            kw[int(c[0])] = dict(star=c[2].count('★'), kw=c[4].split('\n')[0].strip())
    prefs = []
    for line in PREF_TXT.split('\n'):
        n, ja, ko, cja, cko, r = line.split('|'); n = int(n)
        stem = ja if ja == '北海道' else ja[:-1]
        prefs.append(dict(n=n, ja=ja, ko=ko, cja=cja, cko=cko, r=int(r), diff=cja != stem,
                          star=kw.get(n, {}).get('star', 1), kw=kw.get(n, {}).get('kw', '')))
    return prefs

def build_map():
    pts = jm.pref_points(jm.get_data()); K = 40; R = 0.812
    proj = lambda lon, lat: np.array(((lon-128.4)*K*R, (45.7-lat)*K))
    P = [np.array([proj(*q) for q in np.array(p, dtype=float)]) for p in pts]
    main = np.vstack(P[:46]); mn, mx = main.min(0), main.max(0)
    ok = P[46]; osz = ok.max(0)-ok.min(0); bx, by = mn[0]+10, mn[1]+20; s = 1.8
    P[46] = (ok-ok.min(0))*s + np.array([bx+8, by+8]); box = [bx, by, osz[0]*s+16, osz[1]*s+16]
    pad = 6; ox, oy = mn[0]-pad, mn[1]-pad; Wd = mx[0]-mn[0]+2*pad; H = mx[1]-mn[1]+2*pad
    paths = []
    for xy in P:
        out = [xy[0]]
        for q in xy[1:]:
            if abs(q-out[-1]).sum() > 1.2: out.append(q)
        dstr = 'M'+'L'.join(f'{x-ox:.0f} {y-oy:.0f}' for x, y in out)+'Z'
        c = np.mean(out, axis=0)-[ox, oy]
        paths.append(dict(d=dstr, cx=round(c[0]), cy=round(c[1])))
    return dict(vb=f"0 0 {Wd:.0f} {H:.0f}", inset=[round(box[0]-ox), round(box[1]-oy), round(box[2]), round(box[3])], paths=paths, _off=(float(ox), float(oy)))

YEAR_IN_NAME = re.compile(r'\d{4}')
def build_jphis(d):
    tab = None
    for t in d.tables:
        try: h = [c.text.strip() for c in t.rows[0].cells]
        except Exception: continue
        if h[:2] == ['연도', '사건'] and len(t.rows) > 100: tab = t; break
    hang = re.compile(r'[가-힣]'); rows = []; last = ''
    for r in tab.rows[1:]:
        yr_, ev, core, withc, gr = [x.text.strip() for x in r.cells]
        yr_ = yr_ or last; last = yr_
        ev = ev.replace('\n', ' ')
        ev = re.sub(r'\s*\[~\s*\d{4}\]', '', ev)   # 「日清戦争 [~1895]」 같은 기간 표시는 떼고 시작 연도로
        m = re.search(r'\[(발효|시행예정)\(?(\d{4})\)?\]', ev); eff = int(m.group(2)) if m else None
        ev = re.sub(r'\s*\[(발효|시행예정)\(?\d{4}\)?\]\s*', '', ev).strip()
        li = max((i for i, ch in enumerate(ev) if hang.match(ch)), default=-1)
        ko, ja = ev[:li+1].strip(), ev[li+1:].strip()
        if ja.startswith(')'): ko += ')'; ja = ja[1:].strip()
        m2 = re.match(r'^(\([A-Za-z0-9 .\-/]+\))\s*(.*)$', ja)
        if m2: ko = ko+' '+m2.group(1); ja = m2.group(2)
        rows.append(dict(y=int(yr_), eff=eff, ko=ko, ja=ja, core=core.replace('\n', ' '), w=withc, g=gr.replace('*', '')))
    # 자료집 오류: 1945년 6건이 1942 칸에 병합됨
    for r in rows:
        if r['ko'] in {'오키나와 전투', '히로시마·나가사키 원폭투하', '포츠담 선언', 'GHQ의 간접통치 개시', '노동조합법', '여성참정권 확립'}: r['y'] = 1945
    return [r for r in rows if r['y'] < 2020 and not YEAR_IN_NAME.search(r['ko']+r['ja'])]

def build_world(xlsx):
    ws = openpyxl.load_workbook(xlsx, read_only=True)['연표DB']
    rows = list(ws.iter_rows(min_row=4, values_only=True)); hdr = rows[0]; out = []
    for r in rows[1:]:
        d = dict(zip(hdr, r))
        if not d['사건ID']: continue
        # 2026-10-09 마스터DB v2 기준: S·A·B 전부 + C 중 장기반복도≥4(=C*)만. 일반 C·N 제외, 수록='제외'(자료집 제외 지시)·자료집병합(중복) 행 제외
        g = d['등급']
        if g == 'C' and (d.get('장기반복도') or 0) >= 4: g = 'C*'
        if g not in ('S', 'A', 'B', 'C*'): continue
        if str(d.get('수록') or '').strip() == '제외' or d.get('자료집병합'): continue
        eff = d['연도_발효'] if isinstance(d['연도_발효'], int) and d['연도_발효'] != d['연도_채택'] else None   # '미발효' 같은 글자는 무시
        e = dict(id=d['사건ID'], y=d['연도_채택'], eff=eff, ko=d['사건명'].strip(), ja=(d['사건명_일'] or '').strip(), f=d['분야'],
                 t=(d['유형'] or '').replace('・', '·'), g=g, core=(d['핵심 한 줄'] or '').strip(), c=d['관련국'] or '')
        if str(d.get('일본연표') or '').strip().upper() == 'O': e['jp'] = 1   # 마스터DB 일본연표 표시 → 통합본에서 세계연표 쪽 중복 출제 방지
        m2 = re.match(r'^(\([A-Za-z0-9 .\-/]+\))\s*(.*)$', e['ja'])
        if m2: e['ko'] = e['ko']+' '+m2.group(1); e['ja'] = m2.group(2)
        if e['y'] < 2020 and not YEAR_IN_NAME.search(e['ko']+e['ja']): out.append(e)
    return out

def build_landforms(prefs):
    byja = {p['ja']: p['n'] for p in prefs}
    def P(*names):
        out = []
        for n in names:
            k = [k for k in byja if k.startswith(n)]; assert len(k) == 1, (n, k); out.append(byja[k[0]])
        return out
    items = []
    for line in open(os.path.join(HERE, 'landforms.txt'), encoding='utf-8').read().strip().split('\n'):
        if not line.strip() or line.startswith('#'): continue
        t, ja, ko, pr, note = line.split('|')
        items.append(dict(t=t, ja=ja, ko=ko, p=P(*pr.split(',')) if pr.strip() else [], note=note))
    return items

def core_landform_names(d):
    """자료집 03~06 요약표(구분|명칭|내용, 구분|지형|대표 장소|…)에 실린 이름 = 「진짜 중요한 지형」.
    반환: (토큰 목록, 유형 힌트) — 「飛騨 · 木曽 · 赤石山脈」처럼 접미사가 마지막에만 붙은 것도 풀어서 돌려준다."""
    SUF = ['山脈', '山地', '平野', '盆地', '半島', '海峡', '水道']
    toks = []
    for t in d.tables:
        try: h = [celltext(c._tc) for c in t.rows[0].cells]
        except Exception: continue
        if h[:2] == ['구분', '명칭']: col = 1
        elif h[:3] == ['구분', '지형', '대표 장소']: col = 2
        else: continue
        for r in t.rows[1:]:
            cells = [celltext(c._tc) for c in r.cells]
            kind, txt = cells[0], cells[col]
            inner = re.findall(r'[（(]([^)）]+)[)）]', txt)
            txt = re.sub(r'[（(][^)）]*[)）]', '', txt)
            parts = [x.strip() for x in re.split(r'\s*[·・]\s*', txt) if x.strip()]
            parts += [x.strip() for y in inner for x in re.split(r'\s*[·・]\s*', y) if x.strip()]
            suf = next((sf for sf in SUF if parts and parts[-1].endswith(sf)), '')
            if not suf: suf = {'반도': '半島', '해협': '海峡', '평야': '平野', '분지': '盆地'}.get(kind, '')
            nmain = len(parts) - sum(len(re.split(r'\s*[·・]\s*', y)) for y in inner)
            for i, x in enumerate(parts):
                if i >= nmain: toks.append(('=', x)); continue          # 괄호 안(현 이름·하천 이름)은 정확히 같을 때만
                toks.append(('^', x + suf if suf and not x.endswith(suf) else x))
    return toks

def mark_core(lands, d):
    """「일본 지형」 세트 출제 대상(core). landforms_pick.txt 가 있으면 그 목록만 (2026-10-01 사용자 지정), 없으면 자료집 요약표 기준."""
    pick = os.path.join(HERE, 'landforms_pick.txt')
    if os.path.exists(pick):
        names = [l.split('#')[0].strip() for l in open(pick, encoding='utf-8')]
        names = [n for n in names if n]
        have = {x['ja'] for x in lands}
        miss = [n for n in names if n not in have]; assert not miss, ('landforms_pick.txt 에 없는 지형', miss)
        for x in lands:
            if x['ja'] in names: x['core'] = 1
        return lands
    toks = core_landform_names(d)
    EXTRA = {'琵琶湖', '豊後水道'}   # 요약표 본문에 나오는 이름(淀川 발원지 琵琶湖) · 해협 줄의 「豊後」
    for x in lands:
        ja = x['ja']; stem = re.sub(r'(山脈|山地)$', '', ja)
        hit = ja in EXTRA or any(t == ja or (m == '^' and ((len(t) >= 2 and ja.startswith(t)) or (len(ja) >= 3 and t.startswith(ja))
                                 or (len(stem) >= 2 and t.startswith(stem) and t.endswith(('山脈', '山地')) and ja.endswith(('山脈', '山地'))))) for m, t in toks)
        if x['t'] == '산업·기타': hit = False    # 공업·유산은 「지형」 세트에서 뺌 (도도부현 특징에서 다룸)
        if hit: x['core'] = 1
    return lands

def build_concept():
    out = []
    for f in sorted(glob.glob(os.path.join(HERE, 'concept', '20*.json'))): out += json.load(open(f, encoding='utf-8'))
    return out

if __name__ == '__main__':
    ap = argparse.ArgumentParser(); ap.add_argument('--jp', required=True); ap.add_argument('--db', required=True); ap.add_argument('--wg', help='EJU 세계지리 자료집.docx (자연지리·통계)'); a = ap.parse_args()
    d = docx.Document(a.jp)
    prefs = build_prefs(d)
    jpmap = build_map(); off = jpmap.pop('_off')
    sys.path.insert(0, os.path.join(HERE, 'geo')); from build_geo import build_geo
    geo = build_geo(off)
    lands = mark_core(build_landforms(prefs), d)
    for x in lands:                                   # 지형 지도 문항용 좌표 (geo/landcoords.py)
        if x['ja'] in geo['landxy']: x['xy'] = geo['landxy'][x['ja']]
    data = dict(prefs=prefs, regions=REGIONS, jpmap=jpmap, jphis=build_jphis(d), world=build_world(a.db),
                landforms=lands, concept=build_concept(),
                wmap=geo['wmap'], countries=geo['countries'], climate=geo['climate'],
                meta=dict(built=__import__('datetime').date.today().isoformat()))
    from llpoints import build_ll; data['wll'] = build_ll()   # 세계 지형 위경도 문항 (geo/llpoints.py)
    if a.wg:
        from build_wgeo import build_wgeo
        wg = build_wgeo(a.wg); data.update(wland=wg['wland'], wstat=wg['wstat'], wbelts=wg['belts'])
    out = os.path.join(HERE, '..', 'data.js')
    open(out, 'w', encoding='utf-8').write('window.KITAKU_DATA='+json.dumps(data, ensure_ascii=False, separators=(',', ':'))+';')
    print({k: (len(v) if isinstance(v, list) else '-') for k, v in data.items()}, '->', out)
